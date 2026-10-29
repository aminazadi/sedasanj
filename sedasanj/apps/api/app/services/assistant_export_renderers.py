"""Pure byte renderers for assistant exports; no tenant data is written to disk."""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

from app.errors import ApiError
from app.services.export_cells import safe_cell

MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
    "md": "text/markdown; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def render(fmt: str, columns: dict[str, str], rows: list[dict[str, Any]]) -> bytes:
    if fmt == "json":
        return json.dumps(
            [{key: row.get(key) for key in columns} for row in rows],
            ensure_ascii=False,
            indent=2,
            default=str,
        ).encode("utf-8")
    if fmt == "pdf":
        return render_pdf(columns, rows)
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "خروجی دستیار"
        sheet.sheet_view.rightToLeft = True
        sheet.freeze_panes = "A2"
        sheet.append(list(columns.values()))
        for row in rows:
            values = [ILLEGAL_CHARACTERS_RE.sub("", cell_text(row.get(key))) for key in columns]
            if any(len(value) > 32767 for value in values):
                raise ApiError(
                    "export_too_large",
                    "متن برای سلول اکسل طولانی است؛ قالب JSON یا TXT را انتخاب کنید",
                )
            sheet.append([safe_cell(value) for value in values])
        for cells in sheet:
            for cell in cells:
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                cell.font = Font(name="Vazirmatn", size=11, bold=cell.row == 1)
                # Explicit string cells also protect phone numbers and formula-like values.
                cell.data_type = "s"
        for index in range(1, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 28
        sheet.auto_filter.ref = sheet.dimensions
        output = io.BytesIO()
        book.save(output)
        return output.getvalue()
    buffer = io.StringIO(newline="")
    if fmt == "csv":
        writer = csv.writer(buffer, lineterminator="\r\n")
        writer.writerow(columns.values())
        for row in rows:
            writer.writerow([safe_cell(cell_text(row.get(key))) for key in columns])
        return buffer.getvalue().encode("utf-8-sig")
    if fmt == "md":

        def escaped(value: Any) -> str:
            return (
                cell_text(value)
                .replace("\\", "\\\\")
                .replace("|", "\\|")
                .replace("\r", "")
                .replace("\n", "<br>")
            )

        buffer.write("| " + " | ".join(columns.values()) + " |\n")
        buffer.write("| " + " | ".join("---" for _ in columns) + " |\n")
        for row in rows:
            buffer.write("| " + " | ".join(escaped(row.get(key)) for key in columns) + " |\n")
    elif fmt == "txt":
        for index, row in enumerate(rows, 1):
            buffer.write(f"{index}\n")
            for key, label in columns.items():
                buffer.write(f"{label}: {cell_text(row.get(key))}\n")
            buffer.write("\n")
    else:
        raise ApiError(
            "invalid_request", "قالب خروجی معتبر نیست؛ یکی از قالب‌های مجاز را انتخاب کنید"
        )
    return buffer.getvalue().encode("utf-8")


def render_pdf(columns: dict[str, str], rows: list[dict[str, Any]]) -> bytes:
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas

    font = "AssistantVazirmatn"
    if font not in pdfmetrics.getRegisteredFontNames():
        path = Path(__file__).resolve().parents[1] / "assets/fonts/Vazirmatn-Regular.ttf"
        pdfmetrics.registerFont(TTFont(font, str(path)))
    output = io.BytesIO()
    canvas = Canvas(output, pagesize=A4)
    width, height = A4
    margin, size, leading = 40, 10, 18
    y = height - margin
    canvas.setTitle("خروجی دستیار صداسنج")

    def visual(text: str) -> str:
        return str(get_display(arabic_reshaper.reshape(text), base_dir="R"))

    def line(text: str) -> None:
        nonlocal y
        if y < margin + leading:
            canvas.showPage()
            y = height - margin
        canvas.setFont(font, size)
        canvas.drawRightString(width - margin, y, visual(text))
        y -= leading

    def paragraph(text: str) -> None:
        # Wrap logical text before shaping so each visual line has correct bidi order.
        for part in text.splitlines() or [""]:
            current = ""
            for token in re.findall(r"\S+\s*", part):
                if (
                    pdfmetrics.stringWidth(visual(current + token), font, size)
                    <= width - 2 * margin
                ):
                    current += token
                    continue
                if current:
                    line(current.rstrip())
                    current = ""
                for char in token:
                    if (
                        pdfmetrics.stringWidth(visual(current + char), font, size)
                        > width - 2 * margin
                    ):
                        line(current)
                        current = ""
                    current += char
            line(current.rstrip())

    paragraph("خروجی دستیار صداسنج")
    for index, row in enumerate(rows, 1):
        paragraph(f"ردیف {index}")
        for key, label in columns.items():
            paragraph(f"{label}: {cell_text(row.get(key))}")
        y -= leading / 2
    canvas.save()
    return output.getvalue()
