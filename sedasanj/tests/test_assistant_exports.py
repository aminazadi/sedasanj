from __future__ import annotations

import csv
import io
import json
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.dialects import postgresql

from app import deps
from app.deps import Principal
from app.errors import ApiError, install_error_handlers
from app.models import AssistantExport, Contact
from app.routers import assistant
from app.routers import assistant_exports as routes
from app.schemas import SettingsUpdate
from app.services import assistant_agent, assistant_tools, contacts
from app.services import assistant_exports as exports
from app.services.assistant_export_renderers import render
from app.services.storage import MemoryStorage


class Result:
    def __init__(self, rows):
        self.rows = rows

    def scalars(self):
        return self

    def all(self):
        return self.rows

    def scalar_one_or_none(self):
        return self.rows[0] if self.rows else None

    def __iter__(self):
        return iter(self.rows)


class Session:
    def __init__(self, *batches):
        self.batches = list(batches)
        self.statements = []
        self.added = []
        self.deleted = []

    async def execute(self, stmt):
        self.statements.append(stmt)
        return Result(self.batches.pop(0) if self.batches else [])

    def add(self, row):
        self.added.append(row)

    async def flush(self):
        pass

    async def delete(self, row):
        self.deleted.append(row)

    async def get(self, model, key):
        return None


@pytest.fixture
def principal():
    return Principal(kind="user", id=uuid4(), tenant_id=uuid4(), role="operator")


@pytest.fixture
def storage(monkeypatch):
    store = MemoryStorage()
    monkeypatch.setattr(exports, "get_storage", lambda: store)
    monkeypatch.setattr(routes, "get_storage", lambda: store)
    monkeypatch.setattr(exports.audit, "record", AsyncMock())
    monkeypatch.setattr(exports.ratelimit, "enforce", AsyncMock())
    monkeypatch.setattr(exports.platform, "effective_models", AsyncMock(return_value={}))
    return store


@pytest.mark.parametrize("value", ["=SUM(A1)", "+1", "-1", "@SUM(A1)", "  =1", "\t=1", "\u200e=1"])
def test_csv_and_xlsx_neutralize_formula_cells(value):
    data = render("csv", {"name": "نام"}, [{"name": value}])
    assert data.startswith(b"\xef\xbb\xbf")
    assert list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))[1][0] == "'" + value
    book = load_workbook(io.BytesIO(render("xlsx", {"name": "نام"}, [{"name": value}])))
    assert book.active.sheet_view.rightToLeft
    assert book.active["A2"].data_type == "s"
    assert book.active["A2"].value == "'" + value
    assert contacts._safe_cell(value) == "'" + value


@pytest.mark.parametrize("fmt", ["csv", "json", "md", "txt", "xlsx", "pdf"])
def test_renderers_preserve_persian_and_selected_columns(fmt):
    data = render(
        fmt,
        {"name": "نام", "phone": "شماره"},
        [
            {
                "name": "علی رضایی",
                "phone": "09120000000",
                "hidden": "must-not-export",
            }
        ],
    )
    assert data
    if fmt == "pdf":
        assert data.startswith(b"%PDF-")
        assert b"/FontFile2" in data
    elif fmt == "xlsx":
        sheet = load_workbook(io.BytesIO(data)).active
        assert sheet["A2"].value == "علی رضایی"
        assert sheet["B2"].value == "09120000000"
    else:
        assert "علی رضایی" in data.decode("utf-8-sig")
        assert b"must-not-export" not in data


def test_json_retains_structured_values_and_markdown_escapes_cells():
    rows = [{"text": "الف|ب\nپ", "nested": ["یک", "دو"]}]
    assert json.loads(render("json", {"nested": "داده"}, rows)) == [{"nested": ["یک", "دو"]}]
    assert "الف\\|ب<br>پ" in render("md", {"text": "متن"}, rows).decode()


def test_excel_refuses_silent_text_truncation():
    with pytest.raises(ApiError) as error:
        render("xlsx", {"text": "متن"}, [{"text": "x" * 32768}])
    assert error.value.code == "export_too_large"


def test_column_selection_and_strict_arguments():
    assert exports.selected_columns("contacts", ["city", "name"]) == {"city": "شهر", "name": "نام"}
    for columns in ([], ["name", "name"], ["unknown"]):
        with pytest.raises(ApiError):
            exports.selected_columns("contacts", columns)
    with pytest.raises(ValidationError):
        assistant_tools.ExportArguments.model_validate(
            {"dataset": "contacts", "format": "csv", "tenant_id": str(uuid4())}
        )
    with pytest.raises(ValidationError):
        SettingsUpdate(assistant_export_ttl_minutes=0)
    assert (
        "export_data"
        in SettingsUpdate(assistant_enabled_tools=["export_data"]).assistant_enabled_tools
    )


@pytest.mark.parametrize(
    "dataset,filters",
    [
        ("contacts", {"city": "تهران"}),
        ("contacts", {"province": "unknown"}),
        ("contacts", {"status": "complete"}),
        ("transcripts", {"province": "تهران"}),
        ("follow_up_tasks", {"status": "unknown"}),
        ("calls_analysis", {"from_date": "2026-10-31", "to_date": "2026-10-30"}),
    ],
)
def test_invalid_filters_are_not_silently_ignored(dataset, filters):
    with pytest.raises(ApiError):
        exports.validate_filters(dataset, exports.ExportFilters.model_validate(filters), None)


async def test_create_exports_expiry_audit_metadata_and_storage(monkeypatch, principal, storage):
    monkeypatch.setattr(exports, "load_rows", AsyncMock(return_value=[{"transcript": "سلام"}]))
    session = Session()
    before = datetime.now(UTC)
    result = await exports.create_export(
        session, principal, "transcripts", "json", exports.ExportFilters(), ["transcript"]
    )
    row = session.added[0]
    assert isinstance(row, AssistantExport)
    assert before + timedelta(hours=1) <= row.expires_at <= datetime.now(UTC) + timedelta(hours=1)
    assert row.user_id == principal.id
    assert row.tenant_id == principal.tenant_id
    assert json.loads(storage.objects[row.storage_key]) == [{"transcript": "سلام"}]
    assert result["file"] == result["files"][0]
    assert result["file"]["rows"] == 1
    assert result["file"]["expires_in_minutes"] == 60
    assert "۱ ساعت" in result["notice"]
    assert exports.audit.record.call_args.kwargs["action"] == "assistant.export.create"
    assert assistant_tools.preview(result)["files"] == result["files"]
    run = SimpleNamespace(
        tool_call_id="tool-1",
        tool_name="export_data",
        status="succeeded",
        duration_ms=1,
        result_preview=assistant_tools.preview(result),
        error_code=None,
    )
    assert assistant._tool_run_payload(run)["files"] == result["files"]


async def test_admin_ttl_and_bound_call_reach_loader(monkeypatch, principal, storage):
    loader = AsyncMock(return_value=[])
    monkeypatch.setattr(exports, "load_rows", loader)
    monkeypatch.setattr(
        exports.platform,
        "effective_models",
        AsyncMock(return_value={"assistant_export_ttl_minutes": "15"}),
    )
    call_id, conversation_id = uuid4(), uuid4()
    session = Session()
    result = await exports.create_export(
        session,
        principal,
        "transcripts",
        "txt",
        exports.ExportFilters(),
        bound_call_id=call_id,
        conversation_id=conversation_id,
    )
    assert result["file"]["expires_in_minutes"] == 15
    assert loader.call_args.kwargs["bound_call_id"] == call_id
    assert session.added[0].conversation_id == conversation_id


@pytest.mark.parametrize(
    "dataset,cap",
    [("transcripts", 500), ("calls_analysis", 5000), ("contacts", 5000), ("follow_up_tasks", 5000)],
)
async def test_caps_reject_instead_of_truncating(monkeypatch, principal, storage, dataset, cap):
    if dataset == "contacts":
        principal = Principal(
            kind="user", id=principal.id, tenant_id=principal.tenant_id, role="org_admin"
        )
    monkeypatch.setattr(exports, "load_rows", AsyncMock(return_value=[{}] * (cap + 1)))
    with pytest.raises(ApiError) as error:
        await exports.create_export(Session(), principal, dataset, "json", exports.ExportFilters())
    assert error.value.code == "export_too_large"
    assert "فیلتر" in error.value.message
    assert not storage.objects


async def test_failed_audit_removes_uploaded_object(monkeypatch, principal, storage):
    monkeypatch.setattr(exports, "load_rows", AsyncMock(return_value=[]))
    monkeypatch.setattr(
        exports.audit, "record", AsyncMock(side_effect=RuntimeError("audit unavailable"))
    )
    with pytest.raises(RuntimeError):
        await exports.create_export(
            Session(), principal, "transcripts", "json", exports.ExportFilters()
        )
    assert not storage.objects


async def test_contacts_are_admin_only(principal):
    with pytest.raises(ApiError) as error:
        await exports.load_rows(Session(), principal, "contacts", exports.ExportFilters())
    assert error.value.code == "forbidden"


async def test_contacts_reuse_search_and_canonical_geo(principal):
    admin = Principal(kind="user", id=principal.id, tenant_id=principal.tenant_id, role="org_admin")
    session = Session(
        [Contact(name="علی", phone_number="09120000000", province="تهران", city="تهران")]
    )
    rows = await exports.load_rows(
        session,
        admin,
        "contacts",
        exports.ExportFilters(query="علی", province="استان تهران", city="تهران"),
    )
    assert rows[0]["name"] == "علی"
    compiled = session.statements[0].compile()
    assert "تهران" in compiled.params.values()
    assert "LIMIT" in str(compiled)
    assert "contacts.tenant_id" in str(compiled)
    assert "%علی%" in compiled.params.values()


@pytest.mark.parametrize("dataset", ["calls_analysis", "transcripts", "follow_up_tasks"])
async def test_operator_scope_bound_call_and_filters_in_every_call_dataset(principal, dataset):
    session = Session([])
    bound = uuid4()
    await exports.load_rows(
        session,
        principal,
        dataset,
        exports.ExportFilters(
            query="سلام",
            status="open" if dataset == "follow_up_tasks" else "complete",
            from_date=datetime.now(UTC) - timedelta(days=1),
        ),
        bound_call_id=bound,
    )
    sql = str(session.statements[0].compile(dialect=postgresql.dialect()))
    values = session.statements[0].compile().params
    assert "call_operator_assignments.operator_id" in sql
    assert "call_operator_assignments.superseded_at IS NULL" in sql
    assert "calls.tenant_id" in sql
    assert principal.id in values.values()
    assert bound in values.values()
    assert "LIMIT" in sql
    assert "started_at >=" in sql


async def test_operator_cannot_select_another_operator(principal):
    with pytest.raises(ApiError) as error:
        await exports.load_rows(
            Session(), principal, "transcripts", exports.ExportFilters(operator_id=uuid4())
        )
    assert error.value.code == "forbidden"


async def test_purge_deletes_object_and_row_under_lock(principal, storage):
    row = SimpleNamespace(id=uuid4(), storage_key="exports/expired.json", deleted_at=None)
    await storage.put(row.storage_key, b"sensitive")
    session = Session([row])
    assert await exports.purge_session(session, principal.tenant_id) == 1
    assert row.storage_key not in storage.objects
    assert session.deleted == [row]
    assert row.deleted_at
    sql = str(session.statements[0].compile(dialect=postgresql.dialect()))
    assert "FOR UPDATE SKIP LOCKED" in sql
    assert "assistant_exports.tenant_id" in sql


async def test_cleanup_failure_retains_record_for_retry(monkeypatch, principal, storage):
    row = SimpleNamespace(storage_key="exports/expired.json", deleted_at=None)
    session = Session([row])
    monkeypatch.setattr(
        storage, "delete", AsyncMock(side_effect=RuntimeError("storage unavailable"))
    )
    with pytest.raises(RuntimeError):
        await exports.purge_session(session, principal.tenant_id)
    assert not session.deleted
    assert row.deleted_at is None


async def test_cleanup_visits_tenant_databases(monkeypatch):
    tenants = [uuid4(), uuid4()]
    seen = []

    @asynccontextmanager
    async def scope(tenant):
        seen.append(tenant)
        yield Session()

    monkeypatch.setattr(exports, "export_tenant_ids", AsyncMock(return_value=tenants))
    monkeypatch.setattr(exports, "session_scope", scope)
    monkeypatch.setattr(exports, "purge_session", AsyncMock(return_value=1))
    assert await exports.purge_expired() == 2
    assert seen == tenants


@pytest.mark.parametrize(
    "raw",
    [
        "فایل آماده شد",
        "فایل آماده شد<call-sources>[]</call-sources>",
        "پاسخ<call-sources>oops</call-sources>",
    ],
)
def test_expiry_notice_is_appended_even_when_model_omits_it(raw):
    notice = exports.expiry_notice(datetime.now(UTC) + timedelta(hours=1), 60)
    results = [
        (
            assistant_agent.PlannedToolCall("tool-1", "export_data", {}),
            {"file": {}, "notice": notice},
        )
    ]
    results[0][1]["file"] = {"export_id": str(uuid4())}
    answer, _ = assistant_agent.final_response(raw, set(), results)
    assert notice in answer
    repeated, _ = assistant_agent.final_response(answer, set(), results)
    assert repeated.count(notice) == 1
    assert "notice" in assistant_agent.final_input("خروجی بده", "", results)


@pytest.mark.parametrize(
    "case,expected",
    [
        ("expired", 410),
        ("deleted", 410),
        ("purged", 410),
        ("other_owner", 404),
        ("owner", 200),
        ("admin", 200),
    ],
)
async def test_authenticated_download_enforces_owner_and_expiry(principal, storage, case, expected):
    row = SimpleNamespace(
        id=uuid4(),
        tenant_id=principal.tenant_id,
        user_id=principal.id,
        expires_at=datetime.now(UTC) + timedelta(hours=1),
        deleted_at=None,
        storage_key="exports/test.json",
        format="json",
        filename="test.json",
        size_bytes=2,
    )
    await storage.put(row.storage_key, b"[]")
    if case == "expired":
        row.expires_at = datetime.now(UTC)
    if case == "deleted":
        row.deleted_at = datetime.now(UTC)
    if case in {"other_owner", "admin"}:
        row.user_id = uuid4()
    if case == "admin":
        principal = Principal(
            kind="user", id=principal.id, tenant_id=principal.tenant_id, role="org_admin"
        )
    session = Session([] if case == "purged" else [row])
    app = FastAPI()
    app.include_router(routes.router)
    install_error_handlers(app)
    app.dependency_overrides[deps.current_user] = lambda: principal
    app.dependency_overrides[deps.tenant_session] = lambda: session
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(f"/v1/assistant/exports/{row.id}/download")
    assert response.status_code == expected
    if expected == 200:
        assert response.content == b"[]"
        assert response.headers["cache-control"] == "no-store"
        assert "attachment" in response.headers["content-disposition"]
        assert exports.audit.record.call_args.kwargs["action"] == "assistant.export.download"
    elif expected == 410:
        assert response.json()["error"]["code"] == "export_expired"
    compiled = session.statements[0].compile()
    assert principal.tenant_id in compiled.params.values()


def test_font_license_is_bundled():
    path = Path(exports.__file__).resolve().parents[1] / "assets/fonts"
    assert (path / "Vazirmatn-Regular.ttf").stat().st_size > 10000
    assert "SIL OPEN FONT LICENSE" in (path / "OFL.txt").read_text()


async def test_commit_failure_compensates_uploaded_export(monkeypatch, principal, storage):
    import app.services.storage as storage_module

    export_id = uuid4()
    file = {"export_id": str(export_id), "format": "json"}
    key = f"exports/{principal.tenant_id}/{export_id}.json"
    await storage.put(key, b"[]")

    @asynccontextmanager
    async def failed_commit(tenant_id):
        yield Session()
        raise RuntimeError("commit failed")

    monkeypatch.setattr("app.services.tool_gateway.session_scope", failed_commit)
    monkeypatch.setattr(storage_module, "get_storage", lambda: storage)
    monkeypatch.setattr(assistant_tools, "execute", AsyncMock(return_value=({"file": file}, 1)))
    with pytest.raises(RuntimeError, match="commit failed"):
        await assistant._execute_planned_tool(
            assistant_agent.PlannedToolCall("export-1", "export_data", {}), principal, None
        )
    assert key not in storage.objects


async def test_cleanup_continues_after_one_tenant_fails(monkeypatch):
    tenants = [uuid4(), uuid4()]
    seen = []

    @asynccontextmanager
    async def scope(tenant):
        seen.append(tenant)
        if tenant == tenants[0]:
            raise RuntimeError("tenant unavailable")
        yield Session()

    monkeypatch.setattr(exports, "export_tenant_ids", AsyncMock(return_value=tenants))
    monkeypatch.setattr(exports, "session_scope", scope)
    monkeypatch.setattr(exports, "purge_session", AsyncMock(return_value=1))
    with pytest.raises(RuntimeError, match="1 tenants"):
        await exports.purge_expired()
    assert seen == tenants


async def test_socket_reconnect_preserves_export_cards(monkeypatch, principal):
    from app.services import assistant_socket

    conversation_id, message_id = uuid4(), uuid4()
    file = {"export_id": str(uuid4()), "name": "test.json"}
    message = SimpleNamespace(
        id=message_id,
        role="assistant",
        content="فایل آماده شد",
        status="succeeded",
        model=None,
        sources=[],
        attachments=None,
        created_at=datetime.now(UTC),
    )
    run = SimpleNamespace(
        assistant_message_id=message_id,
        tool_call_id="export-1",
        tool_name="export_data",
        status="succeeded",
        duration_ms=1,
        result_preview={"files": [file]},
        error_code=None,
    )
    conversation = SimpleNamespace(
        id=conversation_id,
        tenant_id=principal.tenant_id,
        owner_id=principal.id,
        mode="normal",
        title="خروجی",
    )
    session = Session([message], [run], [], [])
    session.get = AsyncMock(return_value=conversation)

    @asynccontextmanager
    async def scope(tenant):
        yield session

    monkeypatch.setattr(assistant_socket, "session_scope", scope)
    monkeypatch.setattr(
        assistant_socket.sio,
        "get_session",
        AsyncMock(
            return_value={
                "tenant_id": principal.tenant_id,
                "owner_id": principal.id,
                "role": principal.role,
            }
        ),
    )
    monkeypatch.setattr(assistant_socket.sio, "enter_room", AsyncMock())
    emit = AsyncMock()
    monkeypatch.setattr(assistant_socket.sio, "emit", emit)
    assert await assistant_socket.join_conversation(
        "test-socket", {"conversation_id": str(conversation_id)}
    ) == {"ok": True, "sequence": 0}
    state = emit.call_args.args[1]
    assert state["messages"][0]["tool_runs"][0]["files"] == [file]
    compiled = session.statements[1].compile()
    assert principal.tenant_id in compiled.params.values()
    assert conversation_id in compiled.params.values()


def test_export_help_articles_cover_errors_and_safe_content():
    import importlib.util

    path = Path(__file__).resolve().parents[1] / "scripts/seed_help_articles.py"
    spec = importlib.util.spec_from_file_location("export_help_seed", path)
    assert spec and spec.loader
    seed = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(seed)
    articles = {slug: (codes, html) for slug, _title, codes, html in seed.ARTICLES}
    assert articles["assistant-export-expired"][0] == ["export_expired"]
    assert articles["assistant-export-too-large"][0] == ["export_too_large"]
    assert "assistant-data-export" in articles
    for slug in ["assistant-data-export", "assistant-export-expired", "assistant-export-too-large"]:
        _codes, html = articles[slug]
        for heading in ["علائم", "علت", "راه‌حل", "چطور مطمئن شوم درست شد؟", "پیشگیری"]:
            assert heading in html


async def test_download_requires_authentication():
    app = FastAPI()
    app.include_router(routes.router)
    install_error_handlers(app)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.get(f"/v1/assistant/exports/{uuid4()}/download")
    assert response.status_code == 401


async def test_calls_export_preserves_multiple_crm_outcomes_without_duplicate_rows(principal):
    call = SimpleNamespace(
        id=uuid4(),
        started_at=datetime.now(UTC),
        duration_ms=12000,
        caller_number="09120000000",
        dialed_number="101",
        agent_extension="101",
        direction="inbound",
        status="complete",
        call_type="answered",
    )
    insight = SimpleNamespace(
        summary="خلاصه",
        intent="follow_up",
        sentiment="happy",
        topics=["فروش"],
        keywords=["قیمت"],
        action_items=["پیگیری"],
    )
    sale = SimpleNamespace(
        outcome="interested",
        certainty="inferred",
        confidence=0.8,
        product="محصول",
        next_action="تماس",
    )
    outcomes = [
        SimpleNamespace(
            call_id=call.id, outcome=outcome, amount=amount, currency="IRR", product="محصول"
        )
        for outcome, amount in [("interested", None), ("won", 1000)]
    ]
    session = Session([(call, insight, sale)], outcomes)
    rows = await exports.load_rows(session, principal, "calls_analysis", exports.ExportFilters())
    assert len(rows) == 1
    assert rows[0]["sales_inference"]["outcome"] == "interested"
    assert [outcome["outcome"] for outcome in rows[0]["crm_outcomes"]] == ["interested", "won"]
    assert rows[0]["duration_seconds"] == 12
    assert "crm_outcomes.tenant_id" in str(session.statements[1])


async def test_transcript_export_prefers_corrected_text(principal):
    call = SimpleNamespace(
        id=uuid4(),
        started_at=datetime.now(UTC),
        caller_number="09120000000",
        dialed_number="101",
        agent_extension="101",
    )
    transcript = SimpleNamespace(corrected_text="متن اصلاح‌شده", full_text="متن خام")
    session = Session([(call, transcript)])
    rows = await exports.load_rows(session, principal, "transcripts", exports.ExportFilters())
    assert rows[0]["transcript"] == "متن اصلاح‌شده"


async def test_missing_export_migration_is_actionable_and_never_echoes_database_values(
    monkeypatch,
    principal,
    storage,
):
    from sqlalchemy.exc import ProgrammingError

    class MissingTable(Exception):
        sqlstate = "42P01"

    monkeypatch.setattr(exports, "load_rows", AsyncMock(return_value=[]))
    session = Session()
    session.flush = AsyncMock(
        side_effect=ProgrammingError("sensitive query", {}, MissingTable("private detail"))
    )
    with pytest.raises(ApiError) as error:
        await exports.create_export(
            session, principal, "transcripts", "json", exports.ExportFilters()
        )
    assert error.value.code == "assistant_export_unavailable"
    assert "0061_assistant_exports" in error.value.message
    assert "private detail" not in error.value.message
    assert not storage.objects


async def test_export_storage_failure_is_actionable_and_redacted(monkeypatch, principal, storage):
    monkeypatch.setattr(exports, "load_rows", AsyncMock(return_value=[]))
    monkeypatch.setattr(
        storage, "put", AsyncMock(side_effect=RuntimeError("private storage detail"))
    )
    with pytest.raises(ApiError) as error:
        await exports.create_export(
            Session(), principal, "transcripts", "json", exports.ExportFilters()
        )
    assert error.value.code == "assistant_export_unavailable"
    assert "مخزن" in error.value.message
    assert "private storage detail" not in error.value.message
